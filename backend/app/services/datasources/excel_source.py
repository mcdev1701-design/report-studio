"""
Implementacja DataSource dla plików Excel.

Cel:

Odczyt danych z plików XLSX
oraz zwrócenie ich do Report Studio
w postaci Dataset.
"""

from datetime import datetime

from openpyxl import load_workbook

from backend.app.models.dataset import Dataset

from backend.app.services.datasources.file_source import (
    FileSource
)

class ExcelSource(FileSource):
    """
    DataSource dla plików Excel.
    """

    def __init__(
        self,
        file_path: str
    ):
        self.file_path = file_path

        self.connected = False

    def connect(self):
        """
        Walidacja dostępności pliku Excel.
        """

        if not self.test_connection():

            raise FileNotFoundError(
                self.file_path
            )

        self.connected = True

    def disconnect(self):
        """
        Zamknięcie źródła danych.
        """

        self.connected = False

    def test_connection(self):
        """
        Sprawdzenie dostępności pliku.
        """

        try:

            with open(
                self.file_path,
                "rb"
            ):
                pass

            return True

        except FileNotFoundError:

            return False

    def get_data(self) -> Dataset:
        """
        Pobranie danych z pliku Excel.

        Returns:
            Dataset
        """

        if not self.connected:

            raise RuntimeError(
                "DataSource is not connected."
            )

        workbook = load_workbook(
            filename=self.file_path,
            data_only=True
        )

        worksheet = workbook.active

        rows = list(
            worksheet.iter_rows(
                values_only=True
            )
        )

        if not rows:

            return Dataset(
                name="Excel Dataset",
                source_type="excel",
                source_name=self.file_path,
                loaded_at=datetime.now(),
                rows=[]
            )

        headers = rows[0]

        data = []

        for row in rows[1:]:

            data.append(
                dict(
                    zip(
                        headers,
                        row
                    )
                )
            )

        return Dataset(
            name="Excel Dataset",
            source_type="excel",
            source_name=self.file_path,
            loaded_at=datetime.now(),
            rows=data
        )